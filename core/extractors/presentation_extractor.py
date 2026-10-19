"""
core/extractors/presentation_extractor.py

Slide-deck text extraction:
  - .pptx via python-pptx (text from every shape with a text frame,
    plus speaker notes, on every slide)
  - .odp via odfpy
  - .ppt (legacy binary PowerPoint format, pre-2007) is NOT extracted
    - same reasoning as .doc in office_extractor.py: no reliable
      pure-Python reader without a much heavier dependency. Indexed by
      filename only. Documented as a known limitation.
"""

from pathlib import Path
from typing import Optional

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from odf.opendocument import load as odf_load
    from odf.text import P as odf_paragraph
    from odf import teletype as odf_teletype
except ImportError:
    odf_load = None


def extract(path: Path, extension: str) -> Optional[str]:
    ext = extension.lower()
    if ext == ".pptx":
        return _extract_pptx(path)
    if ext == ".odp":
        return _extract_odp(path)
    if ext == ".ppt":
        return None  # legacy binary format - see module docstring
    return None


def _extract_pptx(path: Path) -> Optional[str]:
    if Presentation is None:
        return None
    try:
        presentation = Presentation(str(path))
        parts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text
                if notes.strip():
                    parts.append(notes)
        text = "\n".join(parts)
        return text if text.strip() else None
    except Exception:
        return None


def _extract_odp(path: Path) -> Optional[str]:
    if odf_load is None:
        return None
    try:
        document = odf_load(str(path))
        parts = [
            odf_teletype.extractText(p)
            for p in document.getElementsByType(odf_paragraph)
        ]
        text = "\n".join(p for p in parts if p and p.strip())
        return text if text.strip() else None
    except Exception:
        return None
