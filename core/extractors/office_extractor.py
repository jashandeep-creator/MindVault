"""
core/extractors/office_extractor.py

Word-processor document extraction:
  - .docx via python-docx (paragraphs + table cell text)
  - .odt via odfpy
  - .doc (legacy binary Word format, pre-2007) is NOT extracted - see
    extract_doc() docstring for why. The file is still indexed by
    filename (same pattern as an image with unavailable/failed OCR),
    just without searchable content.
"""

from pathlib import Path
from typing import Optional

try:
    import docx  # python-docx
except ImportError:
    docx = None

try:
    from odf.opendocument import load as odf_load
    from odf.text import P as odf_paragraph
    from odf import teletype as odf_teletype
except ImportError:
    odf_load = None


def extract(path: Path, extension: str) -> Optional[str]:
    ext = extension.lower()
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".odt":
        return _extract_odt(path)
    if ext == ".doc":
        return _extract_doc(path)
    return None


def _extract_docx(path: Path) -> Optional[str]:
    if docx is None:
        return None
    try:
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        text = "\n".join(parts)
        return text if text.strip() else None
    except Exception:
        return None


def _extract_odt(path: Path) -> Optional[str]:
    if odf_load is None:
        return None
    try:
        document = odf_load(str(path))
        parts = [
            odf_teletype.extractText(p)
            for p in document.getElementsByType(odf_paragraph)
        ]
        text = "\n".join(p for p in parts if p and p.strip())
        return text if text.strip() else None
    except Exception:
        return None


def _extract_doc(path: Path) -> Optional[str]:
    """
    Legacy .doc (the pre-2007 binary Word format, OLE Compound File
    based - not a zip like .docx) has no good pure-Python reader
    without a much heavier dependency (e.g. shelling out to
    `antiword`, which isn't pip-installable and isn't present on a
    typical Windows machine). Rather than add a large/unreliable
    dependency for one legacy format, LifeSearch indexes .doc files by
    filename only for now (like a .png with OCR unavailable) - the
    file is still found, just not searchable by content until it's
    converted to .docx. Documented as a known limitation.
    """
    return None
